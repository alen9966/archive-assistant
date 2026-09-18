"""索引 / 清单写出：xlsx 优先，csv 始终（utf-8-sig）。"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

INDEX_COLUMNS = [
    "序号",
    "主文件夹",
    "子文件夹",
    "模板要求文件",
    "实际文件名",
    "新文件名",
    "原文件名",
    "负责人",
    "模板命名规范",
    "项目工号",
    "项目名称",
    "型号",
    "板号",
    "装配变量",
    "套数",
    "版本",
    "日期",
    "来源路径",
    "归档路径",
    "文件哈希",
    "状态",
    "备注",
]


def _row_to_list(row: dict[str, Any]) -> list[str]:
    return ["" if row.get(c) is None else str(row.get(c, "")) for c in INDEX_COLUMNS]


def write_csv(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(INDEX_COLUMNS)
        for row in rows:
            writer.writerow(_row_to_list(row))


def write_xlsx(path: Path, rows: list[dict[str, Any]], sheet_name: str) -> str:
    """成功返回空串，失败返回原因。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "未安装 openpyxl"
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]
        ws.append(INDEX_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(_row_to_list(row))
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        widths = {
            "序号": 8,
            "主文件夹": 22,
            "子文件夹": 28,
            "模板要求文件": 28,
            "实际文件名": 36,
            "新文件名": 42,
            "原文件名": 36,
            "负责人": 22,
            "模板命名规范": 40,
            "来源路径": 40,
            "归档路径": 40,
            "文件哈希": 20,
            "备注": 40,
        }
        for i, col in enumerate(INDEX_COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(INDEX_COLUMNS)):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(path)
        return ""
    except Exception as exc:  # noqa: BLE001
        return str(exc)


def write_table(output_dir: Path, stem: str, rows: list[dict[str, Any]], sheet_name: str) -> dict[str, str]:
    csv_path = output_dir / f"{stem}.csv"
    xlsx_path = output_dir / f"{stem}.xlsx"
    write_csv(csv_path, rows)
    err = write_xlsx(xlsx_path, rows, sheet_name)
    return {"csv": str(csv_path), "xlsx": str(xlsx_path) if not err else "", "xlsx_error": err}
